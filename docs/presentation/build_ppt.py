"""Builds JarCube-myKONE-Demo.pptx. Run: python build_ppt.py (needs python-pptx)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
LOGO = f"{ROOT}/QuantumMind-widget/public/logo.png"
ROBOT = f"{ROOT}/QuantumMind-widget/public/widget-robot.png"
KONE = f"{HERE}/kone-logo.png"
ILL = f"{ROOT}/mykone-portal/packages/design-system/src/assets/illustrations/png"

def rgb(h): return RGBColor.from_string(h.lstrip("#"))
NAVY, BLUE, JBLUE = rgb("13234A"), rgb("1450F5"), rgb("2241FF")
LIGHT, LIGHT2, WHITE = rgb("F5F8FF"), rgb("E7EDFE"), rgb("FFFFFF")
TEXT, GREY, MUTED = rgb("141414"), rgb("5B6478"), rgb("DAE4FF")
RED, GREEN, AMBER, CYAN = rgb("E5391A"), rgb("2EAD5B"), rgb("FFAB00"), rgb("03A6CF")
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=1):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh

def text(s, x, y, w, h, t, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, shape=None, italic=False):
    tb = shape if shape is not None else s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = FONT; r.font.italic = italic
    return tb

def label_box(s, x, y, w, h, t, fill, color=WHITE, size=14, bold=True, line=None,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, align=PP_ALIGN.CENTER):
    sh = box(s, x, y, w, h, fill, line, shape, radius)
    text(s, 0, 0, 0, 0, t.split("\n"), size, color, bold, align, MSO_ANCHOR.MIDDLE, shape=sh)
    return sh

def arrow(s, x1, y1, x2, y2, color=GREY, w=2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c

def header(s, title, sub=None, num=None):
    box(s, 0, 0, 0.18, 7.5, BLUE)
    text(s, 0.6, 0.35, 10.5, 0.7, title, 30, NAVY, True)
    if sub: text(s, 0.6, 1.0, 11, 0.5, sub, 16, GREY)
    s.shapes.add_picture(LOGO, Inches(12.25), Inches(0.3), Inches(0.7), Inches(0.7))
    text(s, 0.6, 7.0, 6, 0.35, "JarCube AI Assistant  ×  myKONE", 10, GREY)
    if num: text(s, 12.2, 7.0, 0.8, 0.35, str(num), 10, GREY, align=PP_ALIGN.RIGHT)

def notes(s, t): s.notes_slide.notes_text_frame.text = t

def bubble(s, x, y, w, h, t, user=False, size=11):
    sh = box(s, x, y, w, h, BLUE if user else LIGHT2, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.25)
    text(s, 0, 0, 0, 0, t, size, WHITE if user else TEXT, shape=sh, anchor=MSO_ANCHOR.MIDDLE)
    return sh

def chat_widget(s, x, y, w, h, msgs, title="myKONE Assistant"):
    """Mock JarCube chat widget. msgs: list of (is_user, text, bubble_height)."""
    box(s, x + 0.05, y + 0.07, w, h, rgb("C9D3EE"), None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    box(s, x, y, w, h, WHITE, rgb("C9D3EE"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    box(s, x, y, w, 0.75, BLUE, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
    box(s, x, y + 0.45, w, 0.3, BLUE)
    s.shapes.add_picture(ROBOT, Inches(x + 0.12), Inches(y + 0.1), Inches(0.62), Inches(0.53))
    text(s, x + 0.8, y + 0.08, w - 1, 0.35, title, 13, WHITE, True)
    text(s, x + 0.8, y + 0.38, w - 1, 0.3, "● Online · powered by JarCube", 9, MUTED)
    cy = y + 0.9
    for user, t, hh in msgs:
        bw = w * 0.78
        bx = x + w - bw - 0.15 if user else x + 0.15
        bubble(s, bx, cy, bw, hh, t, user, 10)
        cy += hh + 0.12
    box(s, x + 0.15, y + h - 0.6, w - 0.3, 0.45, LIGHT, rgb("C9D3EE"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
    text(s, x + 0.25, y + h - 0.57, w - 1, 0.4, "Type your question…", 10, GREY, anchor=MSO_ANCHOR.MIDDLE)
    label_box(s, x + w - 0.6, y + h - 0.56, 0.37, 0.37, "➤", BLUE, size=10, shape=MSO_SHAPE.OVAL)

def portal_frame(s, x, y, w, h, page):
    """Mock myKONE portal window with top navigation."""
    box(s, x, y, w, h, LIGHT, rgb("C9D3EE"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.02)
    box(s, x, y, w, 0.55, NAVY, None)
    s.shapes.add_picture(KONE, Inches(x + 0.2), Inches(y + 0.14), Inches(0.9))
    for i, n in enumerate(["Home", "Equipment", "Tenders", "Reports", "Invoices"]):
        text(s, x + 1.4 + i * 1.15, y + 0.12, 1.1, 0.33, n, 11, WHITE if n == page else MUTED, n == page)
        if n == page: box(s, x + 1.48 + i * 1.15, y + 0.48, 0.8, 0.05, BLUE)

def status_pill(s, x, y, w, h, t, col, size=9.5):
    label_box(s, x, y, w, h, t, col, NAVY if col == AMBER else WHITE, size)

# ================================================================ 1 Title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, 8.4, -1.2, 6.5, 6.5, BLUE, None, MSO_SHAPE.OVAL)
box(s, 10.4, 4.6, 4.5, 4.5, JBLUE, None, MSO_SHAPE.OVAL)
s.shapes.add_picture(ROBOT, Inches(8.9), Inches(1.5), Inches(3.8))
s.shapes.add_picture(LOGO, Inches(0.8), Inches(0.7), Inches(1.1), Inches(1.1))
box(s, 2.1, 0.95, 1.7, 0.62, WHITE, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
s.shapes.add_picture(KONE, Inches(2.25), Inches(1.06), Inches(1.4))
text(s, 0.8, 2.3, 8, 1.8, ["JarCube AI Assistant", "for myKONE"], 48, WHITE, True)
text(s, 0.8, 4.15, 7.3, 1.0, "Instant, secure, 24×7 answers for every myKONE user: tenders, inspection reports and equipment, all in one chat.", 18, MUTED)
box(s, 0.8, 5.45, 2.2, 0.06, WHITE)
text(s, 0.8, 5.6, 7, 0.4, "Product demo  ·  October 2026", 14, MUTED)
notes(s, "Open with the promise: every myKONE user gets an expert assistant that knows their tenders, reports and equipment.")

# ================================================================ 2 Agenda
s = prs.slides.add_slide(BLANK); header(s, "Agenda", num=2)
items = [("01", "Where we are today", "How myKONE users get help right now"),
         ("02", "Meet JarCube", "An AI assistant trained on your platform"),
         ("03", "Use cases", "Tenders · 3PI reports · Equipment · FAQs"),
         ("04", "Safe by design", "The right answer, only to the right user"),
         ("05", "Support team dashboard", "Live inbox, takeover, tickets, insights"),
         ("06", "Any platform, any channel", "Train, plug in, go live"),
         ("07", "Impact & next steps", "Before vs after, and the rollout plan")]
for i, (n, t, d) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 1.5 + row * 1.32
    box(s, x, y, 5.9, 1.12, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.12)
    label_box(s, x + 0.2, y + 0.2, 0.72, 0.72, n, BLUE, size=16, shape=MSO_SHAPE.OVAL)
    text(s, x + 1.1, y + 0.17, 4.7, 0.4, t, 17, NAVY, True)
    text(s, x + 1.1, y + 0.57, 4.7, 0.4, d, 13, GREY)

# ================================================================ 3 Today
s = prs.slides.add_slide(BLANK)
header(s, "Today: getting help on myKONE is slow and manual", "Every question, big or small, goes through people, phones and inboxes", 3)
steps = [("👤", "myKONE user", "has a question"), ("📞", "Calls / emails", "Tender Specialist"),
         ("🎫", "ServiceNow", "ticket raised"), ("🧑‍💻", "Assigned to", "agent / developer"),
         ("⏳", "Waiting…", "hours to days"), ("✉️", "Answer", "by email")]
w, gap, y = 1.75, 0.37, 1.95
for i, (ic, t1, t2) in enumerate(steps):
    x = 0.6 + i * (w + gap)
    box(s, x, y, w, 1.75, rgb("FFF1EE") if i == 4 else LIGHT, rgb("C9D3EE"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    text(s, x, y + 0.15, w, 0.6, ic, 28, NAVY, align=PP_ALIGN.CENTER)
    text(s, x, y + 0.8, w, 0.4, t1, 14, NAVY, True, PP_ALIGN.CENTER)
    text(s, x, y + 1.17, w, 0.4, t2, 12, GREY, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        arrow(s, x + w + 0.03, y + 0.87, x + w + gap - 0.03, y + 0.87, BLUE, 2.5)
pains = [("No self-service", "Even a simple “what is my tender status?” needs a call or an email."),
         ("Slow answers", "Users wait in inboxes and ticket queues while their work stalls."),
         ("Specialists overloaded", "Experts spend their day on repeat questions instead of high-value work."),
         ("Office hours only", "No help in the evening, at weekends or across time zones.")]
for i, (t, d) in enumerate(pains):
    x = 0.6 + i * 3.07
    box(s, x, 4.25, 2.85, 2.4, WHITE, rgb("FFC9BF"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    box(s, x + 0.1, 4.25, 2.65, 0.08, RED)
    text(s, x + 0.15, 4.45, 2.6, 0.45, t, 16, RED, True)
    text(s, x + 0.15, 4.95, 2.6, 1.6, d, 13, TEXT)
notes(s, "Walk through the real journey. Quick queries have no channel at all except email.")

# ================================================================ 4 Meet JarCube
s = prs.slides.add_slide(BLANK)
header(s, "Meet JarCube", "An AI assistant that learns your platform and talks to your users", 4)
s.shapes.add_picture(ROBOT, Inches(0.7), Inches(1.9), Inches(3.6))
text(s, 0.6, 5.2, 4.0, 1.3, "“Ask me anything about your tenders, reports or equipment. I’m here 24×7.”", 15, NAVY, italic=True, align=PP_ALIGN.CENTER)
feats = [("Trained on your data", "Learns from myKONE guides, FAQs, process documents and website content."),
         ("Answers instantly", "Natural conversation at any hour, in seconds instead of days."),
         ("Knows the user", "Answers about the user’s own tenders, reports and equipment."),
         ("Knows its limits", "If it is not sure, it hands over to a specialist instead of guessing."),
         ("Keeps humans in control", "The support team sees every chat and can take over at any moment."),
         ("Works everywhere", "Website widget today; WhatsApp, Telegram and more when needed.")]
for i, (t, d) in enumerate(feats):
    col, row = i % 2, i // 2
    x, y = 4.9 + col * 4.1, 1.65 + row * 1.75
    box(s, x, y, 3.85, 1.55, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    box(s, x, y + 0.2, 0.08, 1.15, BLUE)
    text(s, x + 0.25, y + 0.15, 3.5, 0.4, t, 16, NAVY, True)
    text(s, x + 0.25, y + 0.58, 3.5, 0.95, d, 12.5, GREY)

# ================================================================ 5 How it works
s = prs.slides.add_slide(BLANK)
header(s, "How it works: four simple steps", "Users keep working in myKONE as they do today; the assistant is there on every page", 5)
st = [("1", "Train", "Upload myKONE documents, FAQs and process guides, or point it at the website.", BLUE),
      ("2", "Connect", "Link it to myKONE data: tenders, inspection reports and equipment.", JBLUE),
      ("3", "Embed", "Add the chat widget to the portal once and it shows on every page.", CYAN),
      ("4", "Assist", "Users ask, JarCube answers; anything complex goes to a specialist.", GREEN)]
for i, (n, t, d, c) in enumerate(st):
    x = 0.6 + i * 3.1
    box(s, x, 1.9, 2.8, 3.95, WHITE, rgb("C9D3EE"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    label_box(s, x + 0.95, 2.2, 0.9, 0.9, n, c, size=26, shape=MSO_SHAPE.OVAL)
    text(s, x, 3.3, 2.8, 0.5, t, 22, NAVY, True, PP_ALIGN.CENTER)
    text(s, x + 0.2, 3.9, 2.4, 1.8, d, 13.5, GREY, align=PP_ALIGN.CENTER)
    if i < 3: arrow(s, x + 2.82, 3.85, x + 3.08, 3.85, BLUE, 3)
label_box(s, 0.6, 6.15, 12.1, 0.6, "Ongoing: JarCube collects the questions it could not answer, so the team can teach it and it improves every week.",
          LIGHT2, NAVY, 14, False)

# ================================================================ 6 New journey
s = prs.slides.add_slide(BLANK)
header(s, "Tomorrow: the new support journey", "Most questions are answered on the spot; only real issues reach a specialist", 6)
label_box(s, 0.6, 3.3, 1.9, 1.0, "👤  myKONE user\nasks in chat", NAVY, size=14)
arrow(s, 2.52, 3.8, 3.1, 3.8, BLUE, 2.5)
label_box(s, 3.15, 3.05, 2.3, 1.5, "🤖  JarCube\nunderstands the\nquestion", BLUE, size=14)
arrow(s, 5.47, 3.4, 6.35, 2.1, GREEN, 2.5)
arrow(s, 5.47, 3.8, 6.35, 3.8, GREEN, 2.5)
arrow(s, 5.47, 4.2, 6.35, 5.5, AMBER, 2.5)
label_box(s, 6.4, 1.55, 3.0, 1.1, "📚  General question\nanswered from the\nknowledge base", rgb("E6F6EC"), NAVY, 12.5, line=GREEN)
label_box(s, 6.4, 3.25, 3.0, 1.1, "🔐  Account question\nanswered from the\nuser’s own data", rgb("E6F6EC"), NAVY, 12.5, line=GREEN)
label_box(s, 6.4, 4.95, 3.0, 1.1, "🙋  Not sure / complex\nhanded to a\nTender Specialist", rgb("FFF6E0"), NAVY, 12.5, line=AMBER)
arrow(s, 9.42, 2.1, 10.15, 2.1, GREEN, 2.5); arrow(s, 9.42, 3.8, 10.15, 3.8, GREEN, 2.5)
arrow(s, 9.42, 5.5, 10.15, 5.5, AMBER, 2.5)
label_box(s, 10.2, 1.55, 2.55, 2.8, "✅  Instant answer\n\nin seconds,\n24×7", GREEN, size=15)
label_box(s, 10.2, 4.95, 2.55, 1.1, "🎫  Ticket created\nwith full chat history", AMBER, NAVY, 12.5)
text(s, 0.6, 1.55, 5.0, 1.3, ["No call, no email, no wait.", "Specialists step in only where they add real value, and they already have the full context."], 13, GREY)

# ================================================================ 7 Tender
s = prs.slides.add_slide(BLANK)
header(s, "Use case 1 · “Where is my tender?”", "Customers check tender status themselves, without calling a specialist", 7)
portal_frame(s, 0.6, 1.6, 7.2, 5.2, "Tenders")
text(s, 0.85, 2.3, 4, 0.4, "My tenders", 16, NAVY, True)
rows = [("TND-2041", "Modernisation – Tower A", "In review", AMBER),
        ("TND-2037", "Service contract renewal", "Approved", GREEN),
        ("TND-2029", "Escalator repair – Mall", "Awaiting customer info", RED),
        ("TND-2018", "New elevators – Block C", "Submitted", CYAN)]
box(s, 0.85, 2.8, 6.7, 0.42, LIGHT2)
for j, h in enumerate(["Tender", "Project", "Status"]):
    text(s, 0.95 + [0, 1.4, 4.4][j], 2.83, 2.5, 0.36, h, 11, NAVY, True)
for i, (a, b, c, col) in enumerate(rows):
    y = 3.27 + i * 0.6
    box(s, 0.85, y, 6.7, 0.55, WHITE)
    text(s, 0.95, y + 0.1, 1.4, 0.36, a, 11, BLUE, True)
    text(s, 2.35, y + 0.1, 3.0, 0.36, b, 11, TEXT)
    status_pill(s, 5.3, y + 0.1, 2.1, 0.35, c, col)
text(s, 0.85, 6.35, 6.7, 0.35, "Illustrative mock-up · sample data", 9, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "What is the status of my tender for Tower A?", 0.55),
    (False, "TND-2041 (Modernisation – Tower A) has been In review with the KONE tender team since 28 Sep. Next step: price approval, usually 3–5 working days.", 1.15),
    (True, "Do you need anything from me?", 0.45),
    (False, "Not for this one. But TND-2029 is waiting for your site contact details. Shall I open that form for you?", 0.95)])

# ================================================================ 8 3PI
s = prs.slides.add_slide(BLANK)
header(s, "Use case 2 · Understanding 3PI inspection reports", "Plain-language explanations of findings, severity and next steps", 8)
portal_frame(s, 0.6, 1.6, 7.2, 5.2, "Reports")
text(s, 0.85, 2.3, 6.8, 0.4, "Inspection report · Elevator E-03 · Third-party inspector", 14, NAVY, True)
fnd = [("A-12", "Car door safety contact worn", "Hazardous", RED),
       ("B-07", "Pit lighting below required level", "Needs action", AMBER),
       ("C-03", "Machine room labelling missing", "Minor", CYAN)]
for i, (a, b, c, col) in enumerate(fnd):
    y = 2.85 + i * 0.85
    box(s, 0.85, y, 6.7, 0.72, WHITE, rgb("DCE3F5"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    box(s, 0.85, y + 0.08, 0.1, 0.56, col)
    text(s, 1.05, y + 0.17, 0.8, 0.4, a, 12, NAVY, True)
    text(s, 1.85, y + 0.17, 3.7, 0.4, b, 12, TEXT)
    status_pill(s, 5.65, y + 0.18, 1.7, 0.36, c, col, 10)
label_box(s, 0.85, 5.55, 3.0, 0.5, "💬  Explain this report", BLUE, size=12)
text(s, 0.85, 6.35, 6.7, 0.35, "Illustrative mock-up · sample data", 9, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "Explain my latest inspection report in simple words.", 0.55),
    (False, "There are 3 findings on E-03. One is urgent: A-12, a worn door safety contact, which should be fixed before the next inspection. B-07 and C-03 are smaller fixes.", 1.25),
    (True, "Can KONE fix these?", 0.4),
    (False, "Yes. I can request a repair quote for all 3 findings and share it with your tender specialist.", 0.8)])

# ================================================================ 9 Equipment
s = prs.slides.add_slide(BLANK)
header(s, "Use case 3 · Built into the Equipment page", "The assistant knows which equipment the user is viewing, so there are no IDs to type", 9)
portal_frame(s, 0.6, 1.6, 7.2, 5.2, "Equipment")
text(s, 0.85, 2.3, 6, 0.4, "Elevator E-03 · Tower A, main lobby", 16, NAVY, True)
cards = [("Health", "Good", GREEN), ("Availability", "99.2 %", BLUE),
         ("Open work orders", "2", AMBER), ("Next visit", "14 Oct", NAVY)]
for i, (a, b, c) in enumerate(cards):
    x = 0.85 + (i % 2) * 3.4; y = 2.85 + (i // 2) * 1.3
    box(s, x, y, 3.2, 1.12, WHITE, rgb("DCE3F5"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    text(s, x + 0.15, y + 0.12, 2.9, 0.35, a, 11, GREY)
    text(s, x + 0.15, y + 0.45, 2.9, 0.55, b, 22, c, True)
s.shapes.add_picture(f"{ILL}/site.png", Inches(6.75), Inches(5.4), Inches(0.8))
text(s, 0.85, 6.35, 5.5, 0.35, "Illustrative mock-up · sample data", 9, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "Why was this elevator out of service last week?", 0.55),
    (False, "E-03 was stopped for 3 hours on 29 Sep for a door-operator repair (work order WO-5512). It has been running normally since.", 1.1),
    (True, "When is the next maintenance?", 0.4),
    (False, "A visit is planned for 14 Oct, 09:00–11:00. Do you want me to add it to your calendar?", 0.8)])

# ================================================================ 10 Quick questions
s = prs.slides.add_slide(BLANK)
header(s, "Use case 4 · Quick questions, answered on the spot", "Questions that users email about today get an instant answer instead", 10)
qs = [("How do I add a new user to my company account?", "Portal how-to"),
      ("Which documents do I need to submit a tender?", "Tender process"),
      ("Where can I download my September invoice?", "Invoices"),
      ("What does ‘Awaiting customer info’ mean?", "Status meaning"),
      ("How long does tender approval usually take?", "Timelines"),
      ("Who is my Tender Specialist and how do I reach them?", "Contacts")]
for i, (q, tag) in enumerate(qs):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 1.65 + row * 1.2
    box(s, x, y, 5.9, 1.0, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.15)
    text(s, x + 0.15, y + 0.1, 0.6, 0.8, "💬", 20, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.8, y + 0.1, 3.6, 0.8, q, 13.5, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    label_box(s, x + 4.45, y + 0.3, 1.3, 0.4, tag, WHITE, BLUE, 9.5, line=BLUE)
box(s, 0.6, 5.4, 12.1, 1.35, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
for i, (big, small) in enumerate([("24×7", "always available"), ("Seconds", "instead of days"),
                                  ("Multi-part", "answers every part of a question"), ("Learns", "from unanswered questions")]):
    x = 0.8 + i * 3.0
    text(s, x, 5.5, 2.8, 0.6, big, 24, WHITE, True, PP_ALIGN.CENTER)
    text(s, x, 6.1, 2.8, 0.5, small, 12, MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 11 Safe by design
s = prs.slides.add_slide(BLANK)
header(s, "Safe by design: smart answers, no sensitive leaks", "JarCube is helpful, but only within what each user is allowed to see", 11)
sec = [("🔐", "Only your own data", "Users only get answers about their own company’s tenders, reports and equipment."),
       ("🏢", "Separate data per customer", "Every organisation’s knowledge is stored separately. Nothing crosses over."),
       ("🚫", "Holds back sensitive data", "Internal notes, confidential pricing and other customers’ data are not disclosed."),
       ("🎯", "Grounded, not guessed", "Answers come only from trusted myKONE content. It never makes up policy."),
       ("🙋", "Hands over when unsure", "When confidence is low, the chat goes to a specialist instead of risking a wrong answer."),
       ("👥", "Role-based access", "Admins, agents and auditors each see only what their role allows.")]
for i, (ic, t, d) in enumerate(sec):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.1, 1.7 + row * 2.6
    box(s, x, y, 3.85, 2.35, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    label_box(s, x + 0.25, y + 0.25, 0.75, 0.75, ic, WHITE, size=20, shape=MSO_SHAPE.OVAL, line=BLUE)
    text(s, x + 1.15, y + 0.3, 2.6, 0.7, t, 15, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 1.15, 3.4, 1.15, d, 12.5, GREY)

# ================================================================ 12 Dashboard
s = prs.slides.add_slide(BLANK)
header(s, "The JarCube dashboard for the support team", "One place to watch, join and improve every conversation", 12)
X, Y, W, H = 0.6, 1.6, 8.1, 5.25
box(s, X, Y, W, H, WHITE, rgb("C9D3EE"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.02)
box(s, X, Y, 1.9, H, NAVY)
s.shapes.add_picture(LOGO, Inches(X + 0.15), Inches(Y + 0.15), Inches(0.42), Inches(0.42))
text(s, X + 0.6, Y + 0.18, 1.3, 0.4, "JarCube", 13, WHITE, True)
for i, n in enumerate(["Dashboard", "Conversations", "Service Requests", "Bots", "Question Bank", "Reports", "Channels"]):
    yy = Y + 0.85 + i * 0.5
    if n == "Conversations": box(s, X + 0.1, yy - 0.03, 1.7, 0.4, BLUE, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
    text(s, X + 0.2, yy, 1.65, 0.35, n, 10, WHITE if n == "Conversations" else MUTED)
for i, (a, b, c) in enumerate([("Chats today", "312", BLUE), ("Solved by AI", "81 %", GREEN),
                               ("Handed to team", "59", AMBER), ("Avg. reply", "4 sec", NAVY)]):
    x = X + 2.1 + i * 1.47
    box(s, x, Y + 0.2, 1.37, 0.95, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    text(s, x + 0.08, Y + 0.25, 1.25, 0.3, a, 9, GREY)
    text(s, x + 0.08, Y + 0.52, 1.25, 0.5, b, 18, c, True)
for i, (a, b, c, col) in enumerate([("Acme Facilities", "Tender TND-2041 status?", "AI", GREEN),
                                    ("Nordic Malls", "Explain finding A-12", "AI", GREEN),
                                    ("City Hospital", "Change tender scope", "Agent", AMBER),
                                    ("Harbor Towers", "Invoice copy for Sept", "AI", GREEN)]):
    yy = Y + 1.35 + i * 0.8
    box(s, X + 2.1, yy, 2.75, 0.7, LIGHT2 if i == 2 else LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    text(s, X + 2.2, yy + 0.04, 2.0, 0.3, a, 10, NAVY, True)
    text(s, X + 2.2, yy + 0.34, 2.0, 0.3, b, 9, GREY)
    status_pill(s, X + 4.15, yy + 0.2, 0.62, 0.3, c, col, 8)
px = X + 5.0
box(s, px, Y + 1.35, 2.95, 3.7, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
text(s, px + 0.1, Y + 1.4, 2.8, 0.3, "City Hospital · live", 10, NAVY, True)
bubble(s, px + 0.1, Y + 1.8, 2.3, 0.6, "Can we add 2 escalators to tender TND-2018?", False, 8.5)
bubble(s, px + 0.55, Y + 2.5, 2.3, 0.6, "Connecting you to your Tender Specialist…", True, 8.5)
label_box(s, px + 0.1, Y + 3.25, 2.75, 0.4, "🎫 Ticket SR-1182 created", rgb("FFF6E0"), NAVY, 9, False)
label_box(s, px + 0.1, Y + 4.45, 1.35, 0.45, "Take over", BLUE, size=10)
label_box(s, px + 1.5, Y + 4.45, 1.35, 0.45, "Bot: ON", GREEN, size=10)
text(s, X + 2.1, Y + H - 0.4, 3, 0.3, "Illustrative mock-up", 8, GREY, italic=True)
for i, (t, d) in enumerate([("Live inbox", "Every chat in real time, across all channels"),
                            ("One-click takeover", "A specialist joins and the bot steps back"),
                            ("Service requests", "Tickets with priority, owner and history"),
                            ("Question bank", "Spot unanswered questions and teach the bot"),
                            ("Reports", "Volumes, AI resolution rate, response times")]):
    y = 1.6 + i * 1.07
    box(s, 9.0, y, 3.75, 0.95, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.12)
    text(s, 9.15, y + 0.06, 3.5, 0.35, t, 13.5, NAVY, True)
    text(s, 9.15, y + 0.42, 3.5, 0.5, d, 11, GREY)

# ================================================================ 13 Any platform
s = prs.slides.add_slide(BLANK)
header(s, "One brain, any platform, any channel", "Train once, then plug JarCube into myKONE or any other platform", 13)
text(s, 0.6, 1.55, 3.2, 0.35, "LEARNS FROM", 11, GREY, True)
text(s, 9.5, 1.55, 3.25, 0.35, "TALKS TO USERS ON", 11, GREY, True)
label_box(s, 4.9, 3.15, 3.5, 1.8, "🤖  JarCube\nAI Assistant\n\ntrained on your content", BLUE, size=15)
for i, t in enumerate(["📄  Guides & FAQs", "📋  Process documents", "🌐  Website content", "🔗  Platform data\n(tenders, reports, equipment)"]):
    y = 1.95 + i * 1.2
    label_box(s, 0.6, y, 3.2, 0.95, t, LIGHT2, NAVY, 13)
    arrow(s, 3.82, y + 0.47, 4.88, 4.05, BLUE, 1.75)
for i, t in enumerate(["🖥  myKONE portal widget", "💬  WhatsApp", "✈️  Telegram", "📘  Facebook Messenger", "🧩  Any other web platform"]):
    y = 1.95 + i * 0.97
    label_box(s, 9.5, y, 3.25, 0.8, t, WHITE, NAVY, 13, line=BLUE)
    arrow(s, 8.42, 4.05, 9.48, y + 0.4, BLUE, 1.75)

# ================================================================ 14 Before vs after
s = prs.slides.add_slide(BLANK)
header(s, "Before vs after", "What changes for users, specialists and KONE", 14)
rows = [("Getting help", "Phone call or email", "Chat on any myKONE page"),
        ("Tender status", "Ask a Tender Specialist", "Instant self-service answer"),
        ("3PI reports", "Read alone, or call for help", "Plain-language explanation and next steps"),
        ("Availability", "Office hours", "24×7"),
        ("Response time", "Hours to days", "Seconds (complex cases go to a person)"),
        ("Specialist workload", "Repeat questions all day", "Only complex, high-value cases"),
        ("Tickets", "Raised manually in ServiceNow", "Created with the full chat context"),
        ("Visibility", "Scattered emails and calls", "One dashboard with reports")]
box(s, 3.95, 1.6, 4.3, 0.6, rgb("FFE3DD")); box(s, 8.3, 1.6, 4.45, 0.6, rgb("DDF3E5"))
text(s, 4.05, 1.6, 4.1, 0.6, "❌  Today", 15, RED, True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.4, 1.6, 4.3, 0.6, "✅  With JarCube", 15, GREEN, True, anchor=MSO_ANCHOR.MIDDLE)
for i, (a, b, c) in enumerate(rows):
    y = 2.25 + i * 0.58
    box(s, 0.6, y, 12.15, 0.55, LIGHT if i % 2 == 0 else WHITE)
    text(s, 0.7, y, 3.2, 0.55, a, 13, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.05, y, 4.2, 0.55, b, 12.5, GREY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.4, y, 4.3, 0.55, c, 12.5, TEXT, True, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 15 Rollout
s = prs.slides.add_slide(BLANK)
header(s, "Proposed rollout", "Start small, prove the value, then scale", 15)
box(s, 0.9, 3.05, 11.6, 0.08, LIGHT2)
ph = [("Phase 1", "Pilot", ["Train on myKONE FAQs & guides", "Widget on the portal", "Quick-question answers"], BLUE),
      ("Phase 2", "Connect data", ["Live tender status", "3PI report explanations", "Equipment page assistant"], JBLUE),
      ("Phase 3", "Smart handover", ["Specialist takeover in the dashboard", "Tickets with full chat context", "Hand-off to ServiceNow"], CYAN),
      ("Phase 4", "Scale", ["More countries & languages", "WhatsApp & other channels", "Reuse on other platforms"], GREEN)]
for i, (p, t, bl, c) in enumerate(ph):
    x = 0.6 + i * 3.1
    label_box(s, x + 1.1, 2.75, 0.65, 0.65, str(i + 1), c, size=18, shape=MSO_SHAPE.OVAL)
    text(s, x, 1.65, 2.85, 0.4, p, 13, GREY, True, PP_ALIGN.CENTER)
    text(s, x, 2.0, 2.85, 0.5, t, 20, NAVY, True, PP_ALIGN.CENTER)
    box(s, x, 3.7, 2.85, 2.5, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    box(s, x + 0.1, 3.7, 2.65, 0.08, c)
    text(s, x + 0.15, 3.95, 2.6, 2.2, ["•  " + b for b in bl], 13, TEXT)
label_box(s, 0.6, 6.35, 12.1, 0.5, "Success measures: share of questions solved by AI · response time · specialist hours saved · user satisfaction",
          WHITE, NAVY, 12.5, False, line=BLUE)

# ================================================================ 16 Thank you
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, -1.5, 4.2, 6, 6, BLUE, None, MSO_SHAPE.OVAL)
s.shapes.add_picture(ROBOT, Inches(8.4), Inches(1.7), Inches(4.0))
s.shapes.add_picture(LOGO, Inches(0.8), Inches(0.7), Inches(1.0), Inches(1.0))
text(s, 0.8, 2.1, 7.5, 1.0, "Thank you", 50, WHITE, True)
text(s, 0.8, 3.2, 7.3, 1.2, "From a phone call and a ticket queue to an instant, secure answer on every myKONE page.", 20, MUTED)
text(s, 0.8, 4.6, 7, 0.6, "Questions & live demo", 22, WHITE, True)

out = f"{HERE}/JarCube-myKONE-Demo.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides")
