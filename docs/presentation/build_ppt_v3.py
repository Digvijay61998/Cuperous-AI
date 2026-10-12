"""Builds JarCube-myKONE-Pitch-v3.pptx (Round 2 + DOD + Smart Assist merged).

Run: python build_ppt_v3.py   (needs python-pptx)
Sample data on use-case slides is illustrative; status names and customer tasks
are taken from the myKONE portal translations so they match what customers see.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
LOGO = f"{ROOT}/QuantumMind-widget/public/logo.png"
ROBOT = f"{ROOT}/QuantumMind-widget/public/widget-robot.png"
KONE = f"{HERE}/kone-logo.png"

def rgb(h): return RGBColor.from_string(h.lstrip("#"))
NAVY, BLUE, JBLUE = rgb("13234A"), rgb("1450F5"), rgb("2241FF")
LIGHT, LIGHT2, WHITE = rgb("F5F8FF"), rgb("E7EDFE"), rgb("FFFFFF")
TEXT, GREY, MUTED, BORDER = rgb("141414"), rgb("5B6478"), rgb("DAE4FF"), rgb("C9D3EE")
RED, GREEN, AMBER, CYAN = rgb("E5391A"), rgb("2EAD5B"), rgb("FFAB00"), rgb("03A6CF")
SOFT_RED, SOFT_GREEN, SOFT_AMBER = rgb("FFF1EE"), rgb("E6F6EC"), rgb("FFF6E0")
FONT = "Arial"
FOOTER = "JarCube AI Assistant  ×  myKONE  |  v3"
SAMPLE = "Illustrative mock-up · sample data · statuses and tasks as shown in myKONE"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
_num = [0]


# ------------------------------------------------------------------ helpers
def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=1):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh

def text(s, x, y, w, h, t, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, shape=None, italic=False):
    tb = shape if shape is not None else s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    for i, ln in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = FONT; r.font.italic = italic
    return tb

def label_box(s, x, y, w, h, t, fill, color=WHITE, size=14, bold=True, line=None,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, align=PP_ALIGN.CENTER):
    sh = box(s, x, y, w, h, fill, line, shape, radius)
    text(s, 0, 0, 0, 0, t.split("\n"), size, color, bold, align, MSO_ANCHOR.MIDDLE, shape=sh)
    return sh

def arrow(s, x1, y1, x2, y2, color=GREY, w=2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c

def new_slide(title, sub=None, backup=False):
    s = prs.slides.add_slide(BLANK)
    _num[0] += 1
    box(s, 0, 0, 0.18, 7.5, BLUE)
    text(s, 0.6, 0.35, 10.8, 0.7, title, 26, NAVY, True)
    if sub: text(s, 0.6, 1.0, 11.2, 0.5, sub, 15, GREY)
    s.shapes.add_picture(LOGO, Inches(12.25), Inches(0.3), Inches(0.7), Inches(0.7))
    text(s, 0.6, 7.0, 7, 0.35, FOOTER + ("  |  Backup" if backup else ""), 10, GREY)
    text(s, 12.2, 7.0, 0.8, 0.35, str(_num[0] + 1), 10, GREY, align=PP_ALIGN.RIGHT)
    return s

def notes(s, t): s.notes_slide.notes_text_frame.text = t

def bubble(s, x, y, w, h, t, user=False, size=10):
    sh = box(s, x, y, w, h, BLUE if user else LIGHT2, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
    text(s, 0, 0, 0, 0, t.split("\n"), size, WHITE if user else TEXT, shape=sh, anchor=MSO_ANCHOR.MIDDLE)
    return sh

def chat_widget(s, x, y, w, h, msgs, title="myKONE Assistant"):
    """msgs: list of (is_user, text, bubble_height)."""
    box(s, x + 0.05, y + 0.07, w, h, BORDER, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    box(s, x, y, w, h, WHITE, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
    box(s, x, y, w, 0.75, BLUE, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
    box(s, x, y + 0.45, w, 0.3, BLUE)
    s.shapes.add_picture(ROBOT, Inches(x + 0.12), Inches(y + 0.1), Inches(0.62), Inches(0.53))
    text(s, x + 0.8, y + 0.08, w - 1, 0.35, title, 13, WHITE, True)
    text(s, x + 0.8, y + 0.38, w - 1, 0.3, "● Online · powered by JarCube", 9, MUTED)
    cy = y + 0.9
    for user, t, hh in msgs:
        bw = w * 0.8
        bx = x + w - bw - 0.15 if user else x + 0.15
        bubble(s, bx, cy, bw, hh, t, user, 9.5)
        cy += hh + 0.1
    box(s, x + 0.15, y + h - 0.55, w - 0.3, 0.42, LIGHT, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
    text(s, x + 0.25, y + h - 0.53, w - 1, 0.38, "Ask about your equipment, projects, tenders…", 9, GREY, anchor=MSO_ANCHOR.MIDDLE)
    label_box(s, x + w - 0.55, y + h - 0.51, 0.34, 0.34, "➤", BLUE, size=9, shape=MSO_SHAPE.OVAL)

NAV = ["Home", "Equipment", "Delivery", "Tenders", "Reports", "Invoices"]
def portal_frame(s, x, y, w, h, page):
    box(s, x, y, w, h, LIGHT, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.02)
    box(s, x, y, w, 0.55, NAVY)
    s.shapes.add_picture(KONE, Inches(x + 0.2), Inches(y + 0.14), Inches(0.85))
    for i, n in enumerate(NAV):
        nx = x + 1.25 + i * 0.98
        text(s, nx, y + 0.13, 0.98, 0.33, n, 10, WHITE if n == page else MUTED, n == page)
        if n == page: box(s, nx + 0.08, y + 0.48, 0.7, 0.05, BLUE)

def pill(s, x, y, w, h, t, col, size=9):
    label_box(s, x, y, w, h, t, col, NAVY if col == AMBER else WHITE, size)

def table(s, x, y, widths, header_row, rows, row_h=0.5, size=11, head_fill=NAVY):
    cx = x
    for wdt, h in zip(widths, header_row):
        box(s, cx, y, wdt, row_h, head_fill)
        text(s, cx, y, wdt, row_h, h, size, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
        cx += wdt
    for r, row in enumerate(rows):
        cx, ry = x, y + row_h * (r + 1)
        for c, (wdt, val) in enumerate(zip(widths, row)):
            box(s, cx, ry, wdt, row_h, LIGHT if r % 2 == 0 else WHITE)
            text(s, cx, ry, wdt, row_h, val, size, NAVY if c == 0 else TEXT, c == 0, anchor=MSO_ANCHOR.MIDDLE)
            cx += wdt

def card(s, x, y, w, h, title, body, accent=BLUE, title_size=15, body_size=12, fill=LIGHT):
    box(s, x, y, w, h, fill, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    box(s, x, y + 0.18, 0.08, h - 0.36, accent)
    text(s, x + 0.22, y + 0.12, w - 0.35, 0.45, title, title_size, NAVY, True)
    text(s, x + 0.22, y + 0.58, w - 0.35, h - 0.65, body, body_size, GREY)


# ================================================================ 1 Title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, 8.4, -1.2, 6.5, 6.5, BLUE, None, MSO_SHAPE.OVAL)
box(s, 10.4, 4.6, 4.5, 4.5, JBLUE, None, MSO_SHAPE.OVAL)
s.shapes.add_picture(ROBOT, Inches(8.9), Inches(1.5), Inches(3.8))
s.shapes.add_picture(LOGO, Inches(0.8), Inches(0.7), Inches(1.1), Inches(1.1))
box(s, 2.1, 0.95, 1.7, 0.62, WHITE, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
s.shapes.add_picture(KONE, Inches(2.25), Inches(1.06), Inches(1.4))
text(s, 0.8, 2.2, 8, 1.8, ["myKONE Assistant", "powered by JarCube"], 46, WHITE, True)
text(s, 0.8, 4.05, 7.4, 1.1, "One conversation for every myKONE question: how-to, delivery, equipment, tenders, reports and invoices. It tells customers what is happening, what comes next and what they need to do.", 17, MUTED)
for i, t in enumerate(["Know", "Track", "Act", "Hand over"]):
    label_box(s, 0.8 + i * 1.75, 5.35, 1.6, 0.5, t, WHITE if i == 0 else BLUE, NAVY if i == 0 else WHITE, 14)
text(s, 0.8, 6.2, 7, 0.4, "Pitch v3  ·  October 2026  ·  For myKONE Business Owner, Product Owner & CTO", 13, MUTED)
notes(s, "v3 merges three ideas: Round 2 support assistant, Team DOD actionable delivery agent and myKONE Smart Assist. One assistant, four pillars.")

# ================================================================ 2 Problem
s = new_slide("The problem: the information exists, but is hard to reach",
              "Three teams saw the same gap from three angles")
probs = [("Finding things", "Smart Assist",
          ["Customers don't always know what is in myKONE, where it is or how to use it.",
           "A simple question becomes a search across pages, or a call to KONE."]),
         ("Tracking projects", "Delivery & Operations",
          ["Status is spread over sales orders, screens and long milestone texts.",
           "Customer actions get missed and block the next milestone.",
           "Delivery and service use different flows."]),
         ("Getting help", "Support / Tenders",
          ["Every question needs a call or an email.",
           "Tickets wait hours to days; specialists answer the same questions all day.",
           "Help is only available in office hours."])]
for i, (t, who, pts) in enumerate(probs):
    x = 0.6 + i * 4.1
    box(s, x, 1.7, 3.85, 3.9, WHITE, rgb("FFC9BF"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
    box(s, x + 0.1, 1.7, 3.65, 0.08, RED)
    text(s, x + 0.2, 1.9, 3.5, 0.45, t, 19, RED, True)
    text(s, x + 0.2, 2.35, 3.5, 0.35, "Raised by: " + who, 11, GREY, italic=True)
    text(s, x + 0.2, 2.85, 3.5, 2.7, ["•  " + p for p in pts], 13, TEXT)
label_box(s, 0.6, 5.9, 12.15, 0.85,
          "The result: slower projects, more calls and emails, and specialists spending time on routine questions instead of high-value work.",
          NAVY, WHITE, 15, False)
notes(s, "Credit the internal teams: this is their problem statement. We bring one platform that solves all three.")

# ================================================================ 3 Solution: four pillars
s = new_slide("One assistant on every page, four pillars",
              "A floating chat that knows the portal and knows the customer's own data")
pillars = [("KNOW", "Portal guidance", BLUE,
            ["How-to: users, invoices, reports", "What a status means", "Process & timelines", "Trusted, approved content only"]),
           ("TRACK", "The customer's own data", JBLUE,
            ["Delivery milestones", "Equipment status & visits", "Tenders, reports, invoices", "All equipment in one answer"]),
           ("ACT", "What to do next", CYAN,
            ["Open customer tasks", "Due dates & impact", "One-click link to the right page", "Nothing is missed"]),
           ("HAND OVER", "People when needed", GREEN,
            ["Specialist takes over the chat", "Ticket with full chat history", "Emergencies go straight to KONE", "Bot learns from gaps"])]
for i, (t, sub, c, pts) in enumerate(pillars):
    x = 0.6 + i * 3.08
    box(s, x, 1.7, 2.9, 4.6, WHITE, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
    label_box(s, x, 1.7, 2.9, 0.9, t, c, WHITE, 20, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, x + 0.15, 2.75, 2.6, 0.4, sub, 14, NAVY, True, PP_ALIGN.CENTER)
    text(s, x + 0.2, 3.3, 2.6, 2.9, ["✓  " + p for p in pts], 13, TEXT)
label_box(s, 0.6, 6.45, 12.15, 0.45, "Know = Smart Assist   ·   Track + Act = Team DOD   ·   Know + Track + Hand over = Round 2", LIGHT2, NAVY, 12, False)

# ================================================================ 4 STATUS NEXT ACTION IMPACT
s = new_slide("One pattern: STATUS → NEXT → ACTION → IMPACT",
              "Developed by Team DOD for delivery, reused across the whole portal")
for i, (t, q, c) in enumerate([("STATUS", "Where is it now?", BLUE), ("NEXT", "What happens next?", JBLUE),
                               ("ACTION", "What do I need to do?", CYAN), ("IMPACT", "What if I don't?", AMBER)]):
    x = 0.6 + i * 3.08
    label_box(s, x, 1.6, 2.85, 0.95, f"{t}\n{q}", c, NAVY if c == AMBER else WHITE, 14)
    if i < 3: arrow(s, x + 2.87, 2.07, x + 3.06, 2.07, BLUE, 2.5)
table(s, 0.6, 2.85, [1.7, 2.6, 2.6, 2.7, 2.55],
      ["Area", "Status", "Next", "Action", "Impact"],
      [["Delivery", "Production started", "Shipment, est. 4 Nov", "Confirm installation & handover dates", "Installation slot at risk"],
       ["Equipment", "Out of service, technician on the way", "Visit today 14:00", "None, KONE is on it", "Restored after repair"],
       ["Tenders", "Waiting for PO number", "Work can be scheduled", "Add PO number", "Repair can't be planned"],
       ["Reports", "1 safety issue open", "Repair tender pending", "Review & approve tender", "Finding stays open"],
       ["Invoices", "1 invoice overdue", "Payment reminder", "Pay or contact KONE", "Per contract terms"]],
      row_h=0.6, size=11)
text(s, 0.6, 6.5, 12, 0.4, "Status and next step come from myKONE data. The assistant presents them; it does not guess dates.", 12, GREY, italic=True)

# ================================================================ 5 Delivery
s = new_slide("Use case 1 · Track: “Where is my elevator?”",
              "Delivery status, next milestone and the customer's open tasks in one answer")
portal_frame(s, 0.6, 1.6, 7.2, 5.25, "Delivery")
text(s, 0.85, 2.3, 6.8, 0.4, "Project: Riverside Offices · Elevator 41234561", 14, NAVY, True)
ms = ["Order confirmed", "Drawings approved", "Production started", "Equipment shipped",
      "Delivered to site", "Installation started", "Installation completed", "Final inspection", "Handed over"]
for i, m in enumerate(ms):
    y = 2.8 + i * 0.4
    col = GREEN if i < 2 else (BLUE if i == 2 else BORDER)
    box(s, 1.0, y + 0.07, 0.24, 0.24, col, None, MSO_SHAPE.OVAL)
    if i < len(ms) - 1: box(s, 1.1, y + 0.31, 0.04, 0.16, BORDER)
    text(s, 1.35, y, 3.0, 0.38, m, 11, NAVY if i == 2 else TEXT, i == 2, anchor=MSO_ANCHOR.MIDDLE)
    date = ["12 Jun", "24 Sep", "est. 10 Oct", "est. 4 Nov", "est. 8 Nov", "est. 12 Nov", "", "", ""][i]
    text(s, 4.2, y, 1.2, 0.38, date, 10, GREY, anchor=MSO_ANCHOR.MIDDLE)
box(s, 5.4, 2.8, 2.2, 2.4, WHITE, rgb("FFC9BF"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
text(s, 5.5, 2.85, 2.0, 0.35, "Your tasks", 11, RED, True)
text(s, 5.5, 3.2, 2.05, 2.0, ["• Prepare the shaft for KONE's survey", "• Confirm installation & handover dates", "• Meet the 5 site-readiness requirements"], 9.5, TEXT)
text(s, 0.85, 6.5, 6.8, 0.3, SAMPLE, 8.5, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "Where is my elevator 41234561?", 0.4),
    (False, "STATUS: Production started (drawings approved 24 Sep).\nNEXT: Equipment shipped, est. 4 Nov.\nOn schedule.", 0.95),
    (True, "Do I need to do anything?", 0.4),
    (False, "Yes, 2 tasks before shipment:\n1. Confirm installation and handover dates\n2. Prepare the shaft to KONE's layout\nIMPACT: unconfirmed dates can delay your installation slot.", 1.45),
    (False, "[ Open project in myKONE ]   [ Contact my Project Manager ]", 0.45)])
notes(s, "Milestone names and customer tasks are the real myKONE delivery texts. 'On schedule' and 'action required' are derived from the milestone delayed flag and open tasks.")

# ================================================================ 6 Multi-equipment
s = new_slide("Use case 2 · Track: all my projects in one answer",
              "No switching between sales orders: the assistant summarises every elevator the user can see")
chat_widget(s, 0.6, 1.5, 5.0, 5.3, [
    (True, "Show me my elevator projects.", 0.4),
    (False, "You have 4 active elevators across 2 projects. 1 needs your action.", 0.6),
    (True, "Which one needs action and why?", 0.45),
    (False, "41234563 at Riverside Offices: Delivered to site is next. Please make sure KONE's trucks can access your site on 8 Nov.", 1.0),
    (False, "[ View task ]   [ Remind me later ]", 0.42)])
table(s, 6.0, 1.6, [1.5, 1.8, 2.1, 1.35], ["Equipment", "Project", "Current stage", "Status"],
      [["41234561", "Riverside Offices", "Production started", ""],
       ["41234562", "Riverside Offices", "Drawings approved", ""],
       ["41234563", "Riverside Offices", "Equipment shipped", ""],
       ["41234570", "Harbor Towers", "Installation started", ""]], row_h=0.62, size=11)
for i, (t, c) in enumerate([("On schedule", GREEN), ("On schedule", GREEN), ("Action needed", AMBER), ("Delayed", RED)]):
    pill(s, 10.85, 2.3 + i * 0.62, 1.25, 0.36, t, c, 9)
card(s, 6.0, 4.9, 6.75, 1.65, "Respects myKONE access rights",
     "Only equipment the user is allowed to see in myKONE is included. Users limited to certain sites or contracts only get those sites, and invoices only show with invoice permission.",
     GREEN, 14, 12)
text(s, 6.0, 6.6, 6.7, 0.3, SAMPLE, 8.5, GREY, italic=True)

# ================================================================ 7 Tenders
s = new_slide("Use case 3 · Act: “Where is my tender?”",
              "Customers see what is waiting for them and finish it in one click")
portal_frame(s, 0.6, 1.6, 7.2, 5.25, "Tenders")
text(s, 0.85, 2.3, 4, 0.4, "Active tenders", 15, NAVY, True)
box(s, 0.85, 2.8, 6.7, 0.42, LIGHT2)
for j, h in enumerate(["Tender", "Site", "Status"]):
    text(s, 0.95 + [0, 2.6, 4.6][j], 2.83, 2.5, 0.36, h, 11, NAVY, True)
for i, (a, b, c, col) in enumerate([("Door operator repair", "Tower A", "Pending approval", AMBER),
                                     ("Escalator step chain", "City Mall", "Waiting for PO number", RED),
                                     ("Car lighting upgrade", "Tower B", "Change requested", CYAN),
                                     ("Rope replacement", "Harbor Towers", "Approved", GREEN)]):
    y = 3.27 + i * 0.6
    box(s, 0.85, y, 6.7, 0.55, WHITE)
    text(s, 0.95, y + 0.1, 2.6, 0.36, a, 11, BLUE, True)
    text(s, 3.55, y + 0.1, 2.0, 0.36, b, 11, TEXT)
    pill(s, 5.5, y + 0.1, 1.95, 0.35, c, col)
text(s, 0.85, 6.5, 6.8, 0.3, SAMPLE, 8.5, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "Do any tenders need something from me?", 0.45),
    (False, "Yes, 2 tenders:\n• Door operator repair, Tower A: Pending approval\n• Escalator step chain, City Mall: approved, waiting for your PO number", 1.25),
    (True, "What happens if I don't add the PO?", 0.45),
    (False, "KONE can't schedule the repair until the PO number is added. It takes about a minute.", 0.8),
    (False, "[ Add PO number ]   [ Review Tower A tender ]", 0.42)])
notes(s, "Statuses are the real myKONE tender statuses. Approving and adding a PO happen on the myKONE tender page through a deep link. In-chat actions are a later phase.")

# ================================================================ 8 Reports
s = new_slide("Use case 4 · Inspection reports explained",
              "Plain-language findings, what is a safety issue, and the route to a repair")
portal_frame(s, 0.6, 1.6, 7.2, 5.25, "Reports")
text(s, 0.85, 2.3, 6.8, 0.4, "Inspection report · Elevator 41230018 · Tower A", 14, NAVY, True)
for i, (a, b, c, col) in enumerate([("Car door safety contact worn", "Safety issue", "Repair tender pending", RED),
                                     ("Pit lighting below required level", "Active", "Open", AMBER),
                                     ("Machine room labelling missing", "Active", "Open", CYAN)]):
    y = 2.85 + i * 0.85
    box(s, 0.85, y, 6.7, 0.72, WHITE, rgb("DCE3F5"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    box(s, 0.85, y + 0.08, 0.1, 0.56, col)
    text(s, 1.05, y + 0.17, 3.6, 0.4, a, 11.5, TEXT)
    pill(s, 4.6, y + 0.18, 1.3, 0.36, b, col, 9)
    text(s, 5.95, y + 0.17, 1.6, 0.4, c, 9.5, GREY)
text(s, 0.85, 6.5, 6.8, 0.3, SAMPLE, 8.5, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "Explain my latest inspection report in simple words.", 0.5),
    (False, "3 findings. 1 is a safety issue: the car door safety contact is worn and should be fixed first. KONE has already prepared a repair tender for it. The other 2 are minor.", 1.3),
    (True, "Can KONE fix the other two as well?", 0.45),
    (False, "I'll pass this to your KONE contact with the report attached. A ticket has been created and you'll get an answer by email.", 0.95)])
notes(s, "Customers cannot request a quote themselves in myKONE today, so the bot raises a ticket for the specialist instead of promising a quote.")

# ================================================================ 9 Equipment
s = new_slide("Use case 5 · Equipment status and visits",
              "Context-aware: on the Equipment page the assistant already knows which unit the user is viewing")
portal_frame(s, 0.6, 1.6, 7.2, 5.25, "Equipment")
text(s, 0.85, 2.3, 6.5, 0.4, "Elevator 41230018 · Tower A, main lobby", 15, NAVY, True)
for i, (a, b, c) in enumerate([("Status", "Out of service", RED), ("Open work order", "Callout", AMBER),
                               ("Technician", "On the way", BLUE), ("Last visit", "29 Sep", NAVY)]):
    x = 0.85 + (i % 2) * 3.4; y = 2.85 + (i // 2) * 1.25
    box(s, x, y, 3.2, 1.08, WHITE, rgb("DCE3F5"), MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
    text(s, x + 0.15, y + 0.1, 2.9, 0.35, a, 11, GREY)
    text(s, x + 0.15, y + 0.42, 2.9, 0.55, b, 20, c, True)
steps = ["Created", "Technician assigned", "On the way", "Arrived", "Visit done"]
for i, st in enumerate(steps):
    x = 0.85 + i * 1.34
    pill(s, x, 5.45, 1.25, 0.42, st, BLUE if i == 2 else (GREEN if i < 2 else BORDER), 8.5)
text(s, 0.85, 6.5, 6.8, 0.3, SAMPLE, 8.5, GREY, italic=True)
chat_widget(s, 8.2, 1.4, 4.55, 5.45, [
    (True, "Why is this elevator not working?", 0.42),
    (False, "Elevator 41230018 has been out of service since 10:12 today. A callout is open and the technician is on the way.", 0.95),
    (True, "When was it last serviced?", 0.4),
    (False, "The last maintenance visit was 29 Sep. 2 callouts in the last 90 days.", 0.65),
    (True, "Someone is stuck in the lift!", 0.4),
    (False, "Call KONE's 24/7 service line now: [local number]. Emergency button: press the alarm in the car. I can't handle emergencies in chat.", 0.95)])
notes(s, "Status and work-order steps match the myKONE equipment app. Emergencies are never handled by AI and always get the phone number.")

# ================================================================ 10 Know: quick questions
s = new_slide("Use case 6 · Quick questions, instant answers",
              "Grounded in approved myKONE guides and FAQs. No navigation and no email needed")
qs = [("How do I add a new user to my company account?", "Portal how-to"),
      ("Where can I download my September invoice?", "Invoices"),
      ("What does ‘Waiting for PO number’ mean?", "Status meaning"),
      ("Why can't I see the Invoices page?", "Access & roles"),
      ("What do I need to prepare before installation?", "Delivery process"),
      ("Who is my KONE contact for this site?", "Contacts")]
for i, (q, tag) in enumerate(qs):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 1.65 + row * 1.15
    box(s, x, y, 5.9, 0.98, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.15)
    text(s, x + 0.15, y + 0.1, 0.6, 0.78, "💬", 20, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.8, y + 0.1, 3.6, 0.78, q, 13, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    label_box(s, x + 4.4, y + 0.29, 1.35, 0.4, tag, WHITE, BLUE, 9.5, line=BLUE)
box(s, 0.6, 5.25, 12.15, 1.5, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
for i, (big, small) in enumerate([("24×7", "always available"), ("Seconds", "instead of days"),
                                  ("Approved content", "only, nothing invented"), ("Learns", "from unanswered questions")]):
    x = 0.8 + i * 3.0
    text(s, x, 5.4, 2.8, 0.6, big, 22, WHITE, True, PP_ALIGN.CENTER)
    text(s, x, 6.0, 2.8, 0.5, small, 12, MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 11 Hand over
s = new_slide("Hand over: people step in when they add value",
              "The bot knows its limits. Complex, sensitive or low-confidence questions go to a KONE specialist")
label_box(s, 0.6, 2.9, 2.2, 1.1, "Customer asks", NAVY, size=14)
arrow(s, 2.82, 3.45, 3.4, 3.45, BLUE, 2.5)
label_box(s, 3.45, 2.75, 2.5, 1.4, "Assistant checks:\nconfident? allowed?\nemergency?", BLUE, size=13)
for i, (t, c, fill) in enumerate([("✅  Answer instantly", GREEN, SOFT_GREEN),
                                  ("🙋  Specialist joins the chat live", AMBER, SOFT_AMBER),
                                  ("🎫  Ticket with full chat history", AMBER, SOFT_AMBER),
                                  ("🚨  Emergency: show KONE 24/7 number", RED, SOFT_RED)]):
    y = 1.75 + i * 1.05
    arrow(s, 5.97, 3.45, 6.75, y + 0.42, c, 2)
    label_box(s, 6.8, y, 3.4, 0.85, t, fill, NAVY, 13, line=c)
card(s, 10.45, 1.75, 2.3, 4.0, "Specialists get",
     ["• The whole conversation", "• Customer & equipment context", "• Priority & owner", "• Takeover in one click", "", "Later: tickets synced to ServiceNow"],
     GREEN, 14, 11.5)
label_box(s, 0.6, 6.15, 12.15, 0.6, "Every unanswered question goes to the Question Bank. The team teaches the bot and it improves every week.",
          LIGHT2, NAVY, 13, False)

# ================================================================ 12 Screenshot help
s = new_slide("Coming next · Screenshot help",
              "From Smart Assist: users upload a screenshot of what confuses them and get guidance for that exact screen")
chat_widget(s, 0.6, 1.5, 5.0, 5.3, [
    (True, "[ screenshot: Tender page with greyed-out Approve button ]\nWhy can't I approve this?", 0.85),
    (False, "Approving repair tenders needs the 'approve repair tenders' permission. Your role doesn't include it.", 0.85),
    (False, "Ask your company's myKONE admin to update your role, or I can send the request to your KONE contact.", 0.85)])
for i, (t, d) in enumerate([("Context-aware", "Combines the screenshot, the page the user is on and the trusted knowledge base."),
                            ("Controlled", "Answers stay inside the approved myKONE knowledge. The image is used only for that answer."),
                            ("Private by default", "Screenshots are stored securely, with a retention period agreed with KONE."),
                            ("Status", "Planned for phase 4. Needs a vision-capable AI model, for example on Amazon Bedrock as proposed by Smart Assist.")]):
    card(s, 6.0, 1.5 + i * 1.33, 6.75, 1.18, t, d, CYAN if i < 3 else AMBER, 14, 11.5)
notes(s, "Be honest: this is on the roadmap, not in the live demo.")

# ================================================================ 13 How it works (business level)
s = new_slide("How it works", "KONE stays in control of its data. JarCube handles conversation and knowledge")
label_box(s, 0.6, 2.9, 2.2, 1.3, "👤 myKONE user\nsigned in to\nthe portal", NAVY, size=13)
arrow(s, 2.82, 3.55, 3.5, 3.55, BLUE, 2.5)
label_box(s, 3.55, 2.65, 2.6, 1.8, "🤖 JarCube\nAssistant\n\nunderstands the question\nand chooses a source", BLUE, size=13)
arrow(s, 6.17, 3.15, 7.0, 2.25, BLUE, 2); arrow(s, 6.17, 3.55, 7.0, 3.95, BLUE, 2); arrow(s, 6.17, 3.95, 7.0, 5.6, AMBER, 2)
label_box(s, 7.05, 1.6, 2.6, 1.3, "📚 Knowledge base\nguides, FAQs,\nstatus glossary", LIGHT2, NAVY, 12)
label_box(s, 7.05, 3.3, 2.6, 1.3, "🔐 myKONE Assistant\nconnector\n(runs at KONE)", SOFT_GREEN, NAVY, 12, line=GREEN)
label_box(s, 7.05, 5.0, 2.6, 1.2, "🙋 KONE specialists\ninbox & tickets", SOFT_AMBER, NAVY, 12, line=AMBER)
arrow(s, 9.67, 3.95, 10.25, 3.95, GREEN, 2)
box(s, 10.3, 1.6, 2.45, 4.6, WHITE, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.05)
text(s, 10.4, 1.7, 2.3, 0.35, "myKONE services", 12, NAVY, True, PP_ALIGN.CENTER)
for i, n in enumerate(["Delivery", "Equipment & work orders", "Tenders", "Reports", "Invoices"]):
    label_box(s, 10.45, 2.15 + i * 0.78, 2.15, 0.62, n, LIGHT, NAVY, 11, line=BORDER)
text(s, 0.6, 4.75, 6.2, 1.9, ["What the connector guarantees:",
                              "•  Uses the user's own myKONE sign-in, so they see exactly what they see in the portal",
                              "•  Applies myKONE roles, permissions and site/contract restrictions",
                              "•  Customer data stays at KONE. Only the answer goes back to the chat"], 12, TEXT)
notes(s, "Business-level only. Technical detail: short-lived assistant token issued by the portal, connector as a KONE microapp using existing token verification and permissions.")

# ================================================================ 14 Safe by design
s = new_slide("Safe by design", "Helpful, but only within what each user is allowed to see")
sec = [("🔐", "Same rights as the portal", "The assistant uses the user's myKONE sign-in and permissions. No extra access is created."),
       ("🏢", "Separate data per customer", "Each customer's conversations and data are kept apart. Nothing crosses over."),
       ("🎯", "Grounded, not guessed", "Dates and statuses come from myKONE; policy answers come from approved content only."),
       ("🙋", "Hands over when unsure", "Low confidence, sensitive topics and emergencies go to people, not to the AI."),
       ("🚫", "Holds back sensitive data", "No internal notes, cost prices or other customers' data in answers."),
       ("📋", "Auditable", "Every answer can be traced to its source; KONE staff impersonation is flagged.")]
for i, (ic, t, d) in enumerate(sec):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.1, 1.65 + row * 2.6
    box(s, x, y, 3.85, 2.35, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    label_box(s, x + 0.25, y + 0.25, 0.75, 0.75, ic, WHITE, size=20, shape=MSO_SHAPE.OVAL, line=BLUE)
    text(s, x + 1.15, y + 0.3, 2.6, 0.7, t, 15, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 1.15, 3.4, 1.15, d, 12, GREY)

# ================================================================ 15 Training & quality
s = new_slide("How we train it and prove it is right",
              "Training means curated content plus measured quality, not uploading everything")
steps = [("1", "Knowledge pack", "myKONE guides, FAQs, Service Transparency pages and process docs, rewritten into short, plain-language articles.", BLUE),
         ("2", "Status glossary", "Every status, milestone and customer task: what it means, what's next, what to do. This powers the NEXT and ACTION answers.", JBLUE),
         ("3", "Test set", "100+ real customer questions with approved answers. We measure accuracy before every release.", CYAN),
         ("4", "Learn loop", "Unanswered questions are reviewed weekly, the knowledge is extended and the test set grows.", GREEN)]
for i, (n, t, d, c) in enumerate(steps):
    x = 0.6 + i * 3.1
    box(s, x, 1.75, 2.85, 3.6, WHITE, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    label_box(s, x + 1.0, 2.0, 0.85, 0.85, n, c, size=24, shape=MSO_SHAPE.OVAL)
    text(s, x, 3.0, 2.85, 0.5, t, 18, NAVY, True, PP_ALIGN.CENTER)
    text(s, x + 0.2, 3.5, 2.45, 1.8, d, 12, GREY, align=PP_ALIGN.CENTER)
    if i < 3: arrow(s, x + 2.87, 3.55, x + 3.08, 3.55, BLUE, 2.5)
box(s, 0.6, 5.65, 12.15, 1.15, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.1)
for i, (big, small) in enumerate([("Accuracy", "answers match approved content"), ("Groundedness", "no invented facts"),
                                  ("Right hand-over", "declines what it doesn't know"), ("Languages", "start EN, add by country")]):
    x = 0.8 + i * 3.0
    text(s, x, 5.7, 2.8, 0.5, big, 18, WHITE, True, PP_ALIGN.CENTER)
    text(s, x, 6.2, 2.8, 0.5, small, 11, MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 16 Dashboard
s = new_slide("The support team's view", "One dashboard to watch, join and improve every conversation")
for i, (t, d) in enumerate([("Live inbox", "Every chat in real time, across channels"),
                            ("One-click takeover", "A specialist joins and the bot steps back"),
                            ("Tickets", "Priority, owner and full chat history"),
                            ("Question bank", "See unanswered questions and teach the bot"),
                            ("Reports", "Volumes, AI-solved share, response times, satisfaction")]):
    card(s, 0.6 + (i % 3) * 4.1, 1.7 + (i // 3) * 2.15, 3.85, 1.9, t, d, BLUE, 16, 12.5)
card(s, 8.8, 3.85, 3.95, 1.9, "Channels", "myKONE widget first. The same assistant can later serve WhatsApp, Teams or other KONE platforms.", GREEN, 16, 12.5)
label_box(s, 0.6, 6.05, 12.15, 0.6, "KONE decides which questions the bot may answer and which always go to a person.",
          LIGHT2, NAVY, 13, False)

# ================================================================ 17 Impact
s = new_slide("Impact", "Measured from day one of the pilot")
table(s, 0.6, 1.7, [2.6, 3.5, 3.5], ["", "Today", "With the assistant"],
      [["Getting help", "Phone call or email", "Chat on any myKONE page"],
       ["Delivery status", "Several screens and long texts", "One answer with next step & action"],
       ["Customer tasks", "Easy to miss", "Shown with due date and impact"],
       ["Tender status", "Ask a specialist", "Self-service, with a link to act"],
       ["Response time", "Hours to days", "Seconds; complex cases go to people"],
       ["Availability", "Office hours", "24×7"]], row_h=0.62, size=12)
box(s, 10.45, 1.7, 2.3, 4.34, NAVY, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
text(s, 10.5, 1.85, 2.2, 0.4, "Example estimate", 12, MUTED, True, PP_ALIGN.CENTER)
text(s, 10.5, 2.3, 2.2, 0.7, "200", 34, WHITE, True, PP_ALIGN.CENTER)
text(s, 10.5, 2.95, 2.2, 0.4, "queries / month", 11, MUTED, align=PP_ALIGN.CENTER)
text(s, 10.5, 3.4, 2.2, 0.4, "× 10 min saved", 13, WHITE, True, PP_ALIGN.CENTER)
text(s, 10.5, 3.95, 2.2, 0.7, "≈ 33 h", 34, WHITE, True, PP_ALIGN.CENTER)
text(s, 10.5, 4.6, 2.2, 1.3, "saved per month, before counting fewer delays and fewer specialist emails", 10.5, MUTED, align=PP_ALIGN.CENTER)
label_box(s, 0.6, 6.15, 12.15, 0.6, "KPIs: AI-solved share · response time · specialist hours saved · customer tasks on time · satisfaction",
          WHITE, NAVY, 13, False, line=BLUE)
notes(s, "The 10 min/query figure is from the Smart Assist proposal. Replace with pilot data.")

# ================================================================ 18 Rollout
s = new_slide("Rollout: start small, prove value, scale", "One pillar at a time; every phase is usable on its own")
box(s, 0.9, 3.05, 11.6, 0.08, LIGHT2)
ph = [("Phase 1", "Know", ["Assistant on every myKONE page", "Trained on guides, FAQs & status glossary", "Hand-over and tickets in place"], BLUE),
      ("Phase 2", "Track & Act", ["Delivery: STATUS → NEXT → ACTION → IMPACT", "Tenders and inspection reports", "Equipment status & visits"], JBLUE),
      ("Phase 3", "Full portal", ["Invoices", "Page-aware answers (no IDs to type)", "ServiceNow ticket hand-off"], CYAN),
      ("Phase 4", "Scale", ["Screenshot help", "More countries & languages", "More channels & KONE platforms"], GREEN)]
for i, (p, t, bl, c) in enumerate(ph):
    x = 0.6 + i * 3.1
    label_box(s, x + 1.1, 2.75, 0.65, 0.65, str(i + 1), c, size=18, shape=MSO_SHAPE.OVAL)
    text(s, x, 1.65, 2.85, 0.4, p, 13, GREY, True, PP_ALIGN.CENTER)
    text(s, x, 2.0, 2.85, 0.5, t, 20, NAVY, True, PP_ALIGN.CENTER)
    box(s, x, 3.7, 2.85, 2.4, LIGHT, None, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    box(s, x + 0.1, 3.7, 2.65, 0.08, c)
    text(s, x + 0.15, 3.95, 2.6, 2.1, ["•  " + b for b in bl], 12.5, TEXT)
label_box(s, 0.6, 6.3, 12.15, 0.5, "Pilot proposal: 1 country · 1 customer group · Phase 1 + delivery tracking · 8 weeks · go/no-go on the KPIs",
          NAVY, WHITE, 13, False)

# ================================================================ 19 What we need
s = new_slide("What we need from KONE to start the pilot", "Decisions and access, no big programme")
asks = [("Business Owner", ["Pilot country and customer group", "Agree success KPIs and targets", "Which topics always go to a person"]),
        ("Product Owner", ["Approved content: guides, FAQs, Service Transparency", "Status and task wording sign-off", "Placement of the assistant in the portal"]),
        ("CTO / Architecture", ["Test account and test data", "Approval for the KONE-hosted connector", "Security & data-privacy review, AI model choice"])]
for i, (who, pts) in enumerate(asks):
    x = 0.6 + i * 4.1
    box(s, x, 1.7, 3.85, 4.2, WHITE, BORDER, MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
    label_box(s, x, 1.7, 3.85, 0.75, who, BLUE, WHITE, 17, radius=0.1)
    text(s, x + 0.25, 2.7, 3.4, 3.1, ["✓  " + p for p in pts], 14, TEXT)
label_box(s, 0.6, 6.15, 12.15, 0.6, "Next step: a 30-minute live demo with Phase 1 on a myKONE test environment",
          LIGHT2, NAVY, 14, True)

# ================================================================ 20 Thank you
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY)
box(s, -1.5, 4.2, 6, 6, BLUE, None, MSO_SHAPE.OVAL)
s.shapes.add_picture(ROBOT, Inches(8.4), Inches(1.7), Inches(4.0))
s.shapes.add_picture(LOGO, Inches(0.8), Inches(0.7), Inches(1.0), Inches(1.0))
text(s, 0.8, 2.1, 7.5, 1.0, "Thank you", 50, WHITE, True)
text(s, 0.8, 3.2, 7.3, 1.4, "Don't just tell customers where their project is. Tell them what happens next and what they need to do to keep it moving.", 19, MUTED)
text(s, 0.8, 4.9, 7, 0.6, "Questions & live demo", 22, WHITE, True)

out = f"{HERE}/JarCube-myKONE-Pitch-v3.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides")
